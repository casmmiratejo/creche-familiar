import streamlit as st
import pandas as pd
from datetime import datetime, timedelta
import os
import io
from io import BytesIO
import qrcode
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import psycopg2
import base64

# Criar pastas base locais para ficheiros temporários (fotos e docs)
os.makedirs('uploads', exist_ok=True)
os.makedirs('relatorios_criancas', exist_ok=True)
os.makedirs('mapas_presencas', exist_ok=True)

# ==========================================
# 1. CONFIGURAÇÃO DA BASE DE DADOS (SUPABASE / POSTGRESQL)
# ==========================================
def get_db_connection():
    if "postgres" in st.secrets:
        db_url = st.secrets["postgres"]["url"]
    else:
        db_url = "postgresql://postgres:password@localhost:5432/postgres"
    return psycopg2.connect(db_url)

def init_db():
    conn = get_db_connection()
    c = conn.cursor()
    
    # Tabela de Utilizadores e Autenticação
    c.execute('''CREATE TABLE IF NOT EXISTS utilizadores (
                    id SERIAL PRIMARY KEY,
                    username TEXT UNIQUE,
                    password TEXT,
                    role TEXT,
                    nome_associado TEXT
                )''')
    
    c.execute("UPDATE utilizadores SET role = 'Educadora' WHERE role = 'Educadora / Secretaria'")
    
    c.execute("SELECT COUNT(*) FROM utilizadores")
    if c.fetchone()[0] == 0:
        c.execute("INSERT INTO utilizadores (username, password, role, nome_associado) VALUES (%s, %s, %s, %s) ON CONFLICT (username) DO NOTHING", 
                  ("admin", "admin123", "Administrador", "Direção"))
        conn.commit()

    # Tabela de Recados e Comunicação
    c.execute('''CREATE TABLE IF NOT EXISTS recados (
                    id SERIAL PRIMARY KEY,
                    data_hora TEXT,
                    remetente TEXT,
                    perfil TEXT,
                    tipo_destino TEXT,
                    alvo_destino TEXT,
                    mensagem TEXT,
                    privado INTEGER DEFAULT 0
                )''')
                
    c.execute('''CREATE TABLE IF NOT EXISTS respostas_recados (
                    id SERIAL PRIMARY KEY,
                    recado_id INTEGER,
                    data_hora TEXT,
                    remetente TEXT,
                    mensagem TEXT
                )''')

    # Tabela Completa de Inscrição da Ama (com coluna de arquivo)
    c.execute('''CREATE TABLE IF NOT EXISTS ficha_ama (
                    id SERIAL PRIMARY KEY,
                    crss TEXT,
                    decreto TEXT,
                    nome_completo TEXT,
                    naturalidade TEXT,
                    data_nascimento TEXT,
                    estado_civil TEXT,
                    morada TEXT,
                    contactos TEXT,
                    sabe_ler_escrever TEXT,
                    habilitacoes_literarias TEXT,
                    habilitacoes_profissionais TEXT,
                    habilitacoes_complementares TEXT,
                    experiencia_anterior TEXT,
                    desempregada TEXT,
                    motivo_desemprego TEXT,
                    motivacao_atividade TEXT,
                    tipo_casa TEXT,
                    num_divisoes INTEGER,
                    regime_casa TEXT,
                    coabitantes TEXT,
                    arquivada INTEGER DEFAULT 0
                )''')

    # Tabela Completa de Inscrição da Criança (com coluna de arquivo)
    c.execute('''CREATE TABLE IF NOT EXISTS ficha_crianca (
                    id SERIAL PRIMARY KEY,
                    processo TEXT,
                    data_entrada TEXT,
                    despacho TEXT,
                    ama_responsavel TEXT,
                    nome_completo TEXT,
                    naturalidade TEXT,
                    data_nascimento TEXT,
                    sexo TEXT,
                    num_identificacao TEXT,
                    validade_identificacao TEXT,
                    nif TEXT,
                    niss TEXT,
                    num_utente_sns TEXT,
                    pai_nome TEXT,
                    pai_morada TEXT,
                    pai_telefone TEXT,
                    pai_telemovel TEXT,
                    pai_email TEXT,
                    pai_profissao TEXT,
                    pai_nif TEXT,
                    pai_local_trabalho TEXT,
                    pai_morada_trabalho TEXT,
                    pai_telefone_trabalho TEXT,
                    pai_horario TEXT,
                    mae_nome TEXT,
                    mae_morada TEXT,
                    mae_telefone TEXT,
                    mae_telemovel TEXT,
                    mae_email TEXT,
                    mae_profissao TEXT,
                    mae_nif TEXT,
                    mae_local_trabalho TEXT,
                    mae_morada_trabalho TEXT,
                    mae_telefone_trabalho TEXT,
                    mae_horario TEXT,
                    a_cargo_de TEXT,
                    responsavel_outro_nome TEXT,
                    responsavel_motivo TEXT,
                    responsavel_parentesco TEXT,
                    agregado_familiar TEXT,
                    total_rendimentos REAL,
                    tipo_casa TEXT,
                    regime_casa TEXT,
                    foto_path TEXT,
                    docs_paths TEXT,
                    arquivada INTEGER DEFAULT 0
                )''')

    # Tabela de Registos Diários / Rotinas
    c.execute('''CREATE TABLE IF NOT EXISTS registos_diarios (
                    id SERIAL PRIMARY KEY,
                    data TEXT,
                    crianca TEXT,
                    ama TEXT,
                    bem_estar TEXT,
                    almoco TEXT,
                    lanche TEXT,
                    sono_manha TEXT,
                    sono_tarde TEXT,
                    fraldas INTEGER,
                    bacio INTEGER,
                    saude_medicacao TEXT,
                    recados TEXT
                )''')
    
    # Tabela de Atividades, Ocorrências e Fotos
    c.execute('''CREATE TABLE IF NOT EXISTS atividades (
                    id SERIAL PRIMARY KEY,
                    data TEXT,
                    crianca TEXT,
                    atividades_sociopedagogicas TEXT,
                    marcos_aprendizagem TEXT,
                    ocorrencias TEXT,
                    observacoes TEXT,
                    fotos_paths TEXT
                )''')
    
    # Tabela de Assiduidade e Presenças (Secretaria)
    c.execute('''CREATE TABLE IF NOT EXISTS assiduidade (
                    id SERIAL PRIMARY KEY,
                    data TEXT,
                    crianca TEXT,
                    hora_entrada TEXT,
                    hora_saida TEXT,
                    notas TEXT
                )''')
    conn.commit()
    c.close()
    conn.close()
    st.cache_data.clear()

init_db()

# ==========================================
# 2. CONFIGURAÇÃO DA INTERFACE E AUTENTICAÇÃO
# ==========================================
st.set_page_config(page_title="Gestão Creche Familiar", page_icon="🏡", layout="wide")

@st.cache_data(ttl=60)
def get_amas(*args, **kwargs):
    conn = get_db_connection()
    try:
        # Verifica se foi pedido para incluir arquivadas
        incluir_arquivadas = kwargs.get('incluir_arquivadas', False)
        if len(args) > 0:
            incluir_arquivadas = args[0]
            
        if incluir_arquivadas:
            df = pd.read_sql_query("SELECT nome_completo FROM ficha_ama", conn)
        else:
            df = pd.read_sql_query("SELECT nome_completo FROM ficha_ama WHERE arquivada = 0", conn)
    except Exception:
        df = pd.read_sql_query("SELECT nome_completo FROM ficha_ama", conn)
    finally:
        conn.close()
        
    if df.empty:
        return []
    return df["nome_completo"].tolist()
    
AMAS = get_amas()


@st.cache_data(ttl=60)
def get_criancas(ama_filtro=None, incluir_arquivadas=False, role="", username=""):
    conn = get_db_connection()
    try:
        # Consulta cirúrgica: traz APENAS o nome, ignorando fotos e PDFs pesados!
        if role == "Pais":
            if incluir_arquivadas:
                df = pd.read_sql_query(
                    "SELECT nome_completo FROM ficha_crianca WHERE pai_nome = %s OR mae_nome = %s", 
                    conn, params=(username, username)
                )
            else:
                df = pd.read_sql_query(
                    "SELECT nome_completo FROM ficha_crianca WHERE (pai_nome = %s OR mae_nome = %s) AND arquivada = 0", 
                    conn, params=(username, username)
                )
        else:
            if ama_filtro:
                if incluir_arquivadas:
                    df = pd.read_sql_query("SELECT nome_completo FROM ficha_crianca WHERE ama_responsavel = %s", conn, params=(ama_filtro,))
                else:
                    df = pd.read_sql_query("SELECT nome_completo FROM ficha_crianca WHERE ama_responsavel = %s AND arquivada = 0", conn, params=(ama_filtro,))
            else:
                if incluir_arquivadas:
                    df = pd.read_sql_query("SELECT nome_completo FROM ficha_crianca", conn)
                else:
                    df = pd.read_sql_query("SELECT nome_completo FROM ficha_crianca WHERE arquivada = 0", conn)
    finally:
        conn.close()
        
    if df.empty:
        return []
    return df["nome_completo"].tolist()



def gerar_qr_code(texto):
    qr = qrcode.QRCode(version=1, box_size=10, border=5)
    qr.add_data(texto)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buffered = BytesIO()
    img.save(buffered, format="PNG")
    return buffered.getvalue()

# Sistema de Sessão e Login
if 'authenticated' not in st.session_state:
    st.session_state.authenticated = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.nome_associado = ""

if not st.session_state.authenticated:
    st.title("🏡 Creche Familiar - Autenticação")
    st.markdown("Por favor, introduza as suas credenciais de acesso para entrar no sistema.")
    
    with st.form("form_login"):
        username_input = st.text_input("Utilizador")
        password_input = st.text_input("Palavra-passe", type="password")
        submit_login = st.form_submit_button("🔑 Entrar", type="primary")
        
        if submit_login:
            conn = get_db_connection()
            c = conn.cursor()
            c.execute("SELECT role, nome_associado FROM utilizadores WHERE username = %s AND password = %s", (username_input, password_input))
            user = c.fetchone()
            c.close()
            conn.close()
            
            if user:
                st.session_state.authenticated = True
                st.session_state.username = username_input
                st.session_state.role = user[0]
                st.session_state.nome_associado = user[1]
                st.success("Autenticação com sucesso! A carregar sistema...")
                st.rerun()
            else:
                st.error("Utilizador ou palavra-passe incorretos.")
    
    st.info("💡 **Problemas no início de sessão? Contacte a educadora responsável**")
    st.stop()

# ==========================================
# 3. BARRA LATERAL E NAVEGAÇÃO SEGURA POR PERFIL
# ==========================================
st.sidebar.title("🏡 Creche Familiar")
st.sidebar.markdown(f"👤 **Utilizador:** {st.session_state.username}")
st.sidebar.markdown(f"🛡️ **Perfil:** {st.session_state.role}")

if st.sidebar.button("🚪 Terminar Sessão"):
    st.session_state.authenticated = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.nome_associado = ""
    st.rerun()

st.sidebar.markdown("---")
st.sidebar.subheader("📌 Navegação")

if st.session_state.role == "Administrador":
    menu = st.sidebar.radio("Menu (Administrador):", [
        "📱 Check-in QR Code (Entrada/Saída)",
        "👤 Ficha da Ama",
        "📄 Ficha da Criança",
        "📝 Rotinas e Cuidados Diários (Ama)", 
        "🎨 Atividades e Ocorrências", 
        "🕒 Secretaria e Assiduidade",
        "📊 Relatórios e Impressão",
        "👩‍🏫 Supervisão (Educadora)",
        "💬 Mural de Recados",
        "👥 Gestão de Utilizadores"
    ])
elif st.session_state.role == "Educadora":
    menu = st.sidebar.radio("Menu (Educadora):", [
        "📱 Check-in QR Code (Entrada/Saída)",
        "👤 Ficha da Ama",
        "📄 Ficha da Criança",
        "📝 Rotinas e Cuidados Diários (Ama)", 
        "🎨 Atividades e Ocorrências", 
        "🕒 Secretaria e Assiduidade",
        "📊 Relatórios e Impressão",
        "👩‍🏫 Supervisão (Educadora)",
        "💬 Mural de Recados"
    ])
elif st.session_state.role == "Ama":
    menu = st.sidebar.radio("Menu (Ama):", [
        "📱 Check-in QR Code (Entrada/Saída)",
        "📝 Rotinas e Cuidados Diários (Ama)", 
        "🎨 Atividades e Ocorrências", 
        "🕒 Secretaria e Assiduidade",
        "💬 Mural de Recados"
    ])
else: # Pais
    menu = st.sidebar.radio("Menu (Pais):", [
        "📱 Check-in QR Code (Entrada/Saída)",
        "👨‍👩‍👧 Espaço dos Pais (Rotinas e Presenças)",
        "💬 Mural de Recados"
    ])

# ==========================================
# 4. MÓDULO: CHECK-IN RÁPIDO QR CODE
# ==========================================
if menu == "📱 Check-in QR Code (Entrada/Saída)":
    st.header("📱 Check-in Rápido de Entrada e Saída")
    st.caption("Acolhimento e entrega de crianças via leitura de QR Code ou registo rápido.")
    
    if st.session_state.role == "Pais":
        crianca_checkin = st.session_state.nome_associado
        st.info(f"👦 Educando associado: **{crianca_checkin}**")
    else:
        # 1. Pede a lista filtrada para o utilizador atual
        lista_criancas = get_criancas(
            role=st.session_state.get('role', ''), 
            username=st.session_state.get('username', '')
        )

# 2. Desenha o selectbox com essa lista
        crianca_checkin = st.selectbox("Selecione a Criança:", lista_criancas)
        
    st.markdown("---")
    
    col_ck1, col_ck2, col_ck3 = st.columns([1, 1, 1])
    
    with col_ck1:
        st.subheader("1. Identificação")
        acompanhante_tipo = st.selectbox("Quem está a entregar / recolher?", 
                                         ["Pai", "Mãe", "Irmão / Irmã", "Avô / Avó", "Outro Familiar Autorizado"],
                                         key="acomp_tipo")
        
        nome_outro_acomp = ""
        if acompanhante_tipo == "Outro Familiar Autorizado":
            nome_outro_acomp = st.text_input("Nome do Acompanhante:")
            
        detalhe_acomp = f"{acompanhante_tipo} ({nome_outro_acomp})" if nome_outro_acomp else acompanhante_tipo


    with col_ck2:
        st.subheader("2. Registo da Hora")
        data_hoje_str = datetime.today().strftime("%Y-%m-%d")
        hora_atual_str = datetime.now().strftime("%H:%M")
        
        st.write(f"📅 **Data:** {datetime.today().strftime('%d/%m/%Y')}")
        st.write(f"⏰ **Hora Atual:** {hora_atual_str}")
        
        # Identificar quem está a fazer a validação na sessão
        quem_valida = f"{st.session_state.username} ({st.session_state.role})"
        
        # 🔒 Restrição: Apenas Admin, Educadora ou Ama podem fazer o registo de presenças
        if st.session_state.role in ["Administrador", "Educadora", "Ama"]:
            c_btn1, c_btn2 = st.columns(2)
            
            if c_btn1.button("🟢 ENTRADA", type="primary", use_container_width=True):
                conn = get_db_connection()
                c = conn.cursor()
                c.execute("SELECT id FROM assiduidade WHERE crianca = %s AND data = %s", (crianca_checkin, data_hoje_str))
                registo_existente = c.fetchone()
                
                # Registo detalhado com quem entregou e quem validou no sistema
                nota_entrada = f"Entregue por: {detalhe_acomp} | Validado por: {quem_valida}"
                
                if registo_existente:
                    c.execute("UPDATE assiduidade SET hora_entrada = %s, notas = %s WHERE id = %s", 
                              (hora_atual_str, nota_entrada, registo_existente[0]))
                else:
                    c.execute("INSERT INTO assiduidade (data, crianca, hora_entrada, hora_saida, notas) VALUES (%s, %s, %s, '', %s)",
                              (data_hoje_str, crianca_checkin, hora_atual_str, nota_entrada))
                conn.commit()
                c.close()
                conn.close()
                st.cache_data.clear()
                st.success(f"Entrada de **{crianca_checkin}** registada às {hora_atual_str} (Validado por {quem_valida})!")
                
            if c_btn2.button("🔴 SAÍDA", use_container_width=True):
                conn = get_db_connection()
                c = conn.cursor()
                c.execute("SELECT id, notas FROM assiduidade WHERE crianca = %s AND data = %s", (crianca_checkin, data_hoje_str))
                registo_existente = c.fetchone()
                
                nota_saida = f"Recolhido por: {detalhe_acomp} | Validado por: {quem_valida}"
                
                if registo_existente:
                    notas_atualizadas = f"{registo_existente[1]} | {nota_saida}" if registo_existente[1] else nota_saida
                    c.execute("UPDATE assiduidade SET hora_saida = %s, notas = %s WHERE id = %s", 
                              (hora_atual_str, notas_atualizadas, registo_existente[0]))
                else:
                    c.execute("INSERT INTO assiduidade (data, crianca, hora_entrada, hora_saida, notas) VALUES (%s, %s, '', %s, %s)",
                              (data_hoje_str, crianca_checkin, hora_atual_str, nota_saida))
                conn.commit()
                c.close()
                conn.close()
                st.cache_data.clear()
                st.warning(f"Saída de **{crianca_checkin}** registada às {hora_atual_str} (Validado por {quem_valida})!")
        else:
            st.info("ℹ️ O registo oficial de entradas e saídas é efetuado exclusivamente pela equipa da creche.")
    
    with col_ck3:
        st.subheader("3. QR Code de Acolhimento")
        
        # Substitua este link pelo URL público exato da sua app no Streamlit Cloud
        url_base_app = "https://creche-familiar.streamlit.app"  # Exemplo de URL
        
        texto_qr = f"{url_base_app}"
        img_qr_bytes = gerar_qr_code(texto_qr)
        st.image(img_qr_bytes, caption="Aponte a câmara para aceder à aplicação", width=160)
# ==========================================
# 5. MÓDULO: FICHA DA AMA (COM CONSULTA E ARQUIVAMENTO)
# ==========================================
elif menu == "👤 Ficha da Ama":
    st.header("👤 Ficha da Ama (Inscrição, Consulta e Arquivamento)")
    st.caption("Centro Regional de Segurança Social (C.R.S.S.) e Decreto-Lei n.º 158/84 de 17 de Maio.")
    
    aba_ama_1, aba_ama_2 = st.tabs(["✨ Nova Inscrição", "🔍 Consultar e Arquivar"])
    
    with aba_ama_1:
        with st.form("form_inscricao_ama_completa"):
            st.subheader("1. Enquadramento Institucional")
            c_i1, c_i2 = st.columns(2)
            crss = c_i1.text_input("C.R.S.S. de Enquadramento")
            decreto = c_i2.text_input("Enquadramento Legal", value="Decreto Lei n.º 158/84 de 17 de Maio")
            
            st.subheader("2. Dados Pessoais da Ama")
            a1, a2, a3, a4 = st.columns(4)
            nome_completo = a1.text_input("Nome Completo da Ama")
            naturalidade = a2.text_input("Naturalidade")
            data_nascimento = a3.date_input("Data de Nascimento", datetime(1985, 1, 1))
            estado_civil = a4.selectbox("Estado Civil", ["Solteiro(a)", "Casado(a)", "Divorciado(a)", "Viúvo(a)", "União de Facto"])
            
            a5, a6, a7 = st.columns(3)
            morada = a5.text_input("Morada Completa")
            contactos = a6.text_input("Contactos Telefónicos / Telemóvel")
            sabe_ler_escrever = a7.selectbox("Sabe ler e escrever?", ["Sim", "Não"])
            
            st.subheader("3. Habilitações e Qualificações")
            habilitacoes_literarias = st.text_input("Habilitações Literárias")
            habilitacoes_profissionais = st.text_area("Habilitações Profissionais")
            habilitacoes_complementares = st.text_area("Habilitações Complementares / Formação Específica")
            
            st.subheader("4. Experiência Anterior e Situação Profissional")
            experiencia_anterior = st.text_area("Experiência anterior com crianças:")
            
            a8, a9 = st.columns(2)
            desempregada = a8.selectbox("Está presentemente desempregada?", ["Não", "Sim"])
            motivo_desemprego = a9.text_input("Qual o motivo? (se aplicável)")
            
            motivacao_atividade = st.text_area("Motivos para exercer a atividade de ama:")
            
            st.subheader("5. Condições da Habitação")
            h1, h2, h3 = st.columns(3)
            tipo_casa = h1.selectbox("Tipo de Casa", ["Vivenda", "Andar", "Outro"])
            num_divisoes = h2.number_input("N.º de Divisões", min_value=1, max_value=20, value=3)
            regime_casa = h3.selectbox("Regime de Habitação", ["Própria", "Alugada"])
            
            st.subheader("6. Coabitantes")
            coabitantes_linhas = []
            for i in range(1, 4):
                co1, co2, co3, co4 = st.columns(4)
                c_nome = co1.text_input(f"Nome {i}", key=f"co_nome_{i}")
                c_parent = co2.text_input(f"Parentesco {i}", key=f"co_parent_{i}")
                c_prof = co3.text_input(f"Profissão {i}", key=f"co_prof_{i}")
                c_cc = co4.text_input(f"CC {i}", key=f"co_cc_{i}")
                if c_nome:
                    coabitantes_linhas.append(f"{c_nome} ({c_parent}) - Prof: {c_prof}")

            submit_ama = st.form_submit_button("💾 Guardar Ficha da Ama")

        # 👉 O IF FICA FORA DO WITH ST.FORM (alinhado com o 'with st.form')
        if submit_ama:
            if nome_completo:
                coabitantes_str = " | ".join(coabitantes_linhas)
                conn = get_db_connection()
                c = conn.cursor()
                c.execute('''INSERT INTO ficha_ama 
                           (crss, decreto, nome_completo, naturalidade, data_nascimento, estado_civil, morada, 
                            contactos, sabe_ler_escrever, habilitacoes_literarias, habilitacoes_profissionais, 
                            habilitacoes_complementares, experiencia_anterior, desempregada, motivo_desemprego, 
                            motivacao_atividade, tipo_casa, num_divisoes, regime_casa, coabitantes, arquivada)
                           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s, 0)''',
                          (crss, decreto, nome_completo, naturalidade, str(data_nascimento), estado_civil, morada,
                           contactos, sabe_ler_escrever, habilitacoes_literarias, habilitacoes_profissionais,
                           habilitacoes_complementares, experiencia_anterior, desempregada, motivo_desemprego,
                           motivacao_atividade, tipo_casa, num_divisoes, regime_casa, coabitantes_str))
                conn.commit()
                c.close()
                conn.close()
                st.cache_data.clear()
                st.success(f"Ficha da ama {nome_completo} guardada com sucesso!")
                st.rerun()
            else:
                st.error("O Nome Completo da Ama é obrigatório.")

    with aba_ama_2:
        st.subheader("🔍 Consultar e Gerir Fichas de Amas")
        todas_amas = get_amas(incluir_arquivadas=True)
        if todas_amas:
            ama_selecionada = st.selectbox("Selecione a Ama para consultar:", todas_amas, key="sel_cons_ama")
            conn = get_db_connection()
            df_detalhe_ama = pd.read_sql_query("SELECT * FROM ficha_ama WHERE nome_completo = %s", conn, params=(ama_selecionada,))
            conn.close()
            
            if not df_detalhe_ama.empty:
                reg_ama = df_detalhe_ama.iloc[0]
                status_arq = "🔴 Arquivada" if reg_ama['arquivada'] == 1 else "🟢 Ativa"
                
                # Cabeçalho limpo com o estado
                col_h1, col_h2 = st.columns([3, 1])
                col_h1.markdown(f"### Ficha de: **{reg_ama['nome_completo']}**")
                col_h2.markdown(f"**Estado:** {status_arq}")
                
                st.markdown("---")
                
                # Exibição organizada por secções
                col_d1, col_d2 = st.columns(2)
                with col_d1:
                    st.markdown("#### 📋 Dados Institucionais e Pessoais")
                    st.write(f"**CRSS:** {reg_ama['crss']}")
                    st.write(f"**Enquadramento Legal:** {reg_ama['decreto']}")
                    st.write(f"**Naturalidade:** {reg_ama['naturalidade']}")
                    st.write(f"**Data de Nascimento:** {reg_ama['data_nascimento']}")
                    st.write(f"**Estado Civil:** {reg_ama['estado_civil']}")
                    st.write(f"**Morada:** {reg_ama['morada']}")
                    st.write(f"**Contactos:** {reg_ama['contactos']}")
                    st.write(f"**Sabe ler e escrever:** {reg_ama['sabe_ler_escrever']}")
                
                with col_d2:
                    st.markdown("#### 🎓 Habilitações e Habitação")
                    st.write(f"**Habilitações Literárias:** {reg_ama['habilitacoes_literarias']}")
                    st.write(f"**Habilitações Profissionais:** {reg_ama['habilitacoes_profissionais']}")
                    st.write(f"**Tipo de Casa:** {reg_ama['tipo_casa']} (Regime: {reg_ama['regime_casa']})")
                    st.write(f"**N.º de Divisões:** {reg_ama['num_divisoes']}")
                    st.write(f"**Coabitantes:** {reg_ama['coabitantes']}")
                
                st.markdown("---")
                st.write(f"**Experiência Anterior:** {reg_ama['experiencia_anterior']}")
                st.write(f"**Motivação para a Atividade:** {reg_ama['motivacao_atividade']}")
                
                st.markdown("---")
                col_arq1, col_arq2 = st.columns(2)
                if reg_ama['arquivada'] == 0:
                    if col_arq1.button("📁 Arquivar Ama", key="btn_arq_ama"):
                        conn = get_db_connection()
                        c = conn.cursor()
                        c.execute("UPDATE ficha_ama SET arquivada = 1 WHERE id = %s", (int(reg_ama['id']),))
                        conn.commit()
                        c.close()
                        conn.close()
                        st.cache_data.clear()
                        st.success("Ama arquivada com sucesso!")
                        st.rerun()
                else:
                    if col_arq2.button("📂 Desarquivar Ama", key="btn_desarq_ama"):
                        conn = get_db_connection()
                        c = conn.cursor()
                        c.execute("UPDATE ficha_ama SET arquivada = 0 WHERE id = %s", (int(reg_ama['id']),))
                        conn.commit()
                        c.close()
                        conn.close()
                        st.cache_data.clear()
                        st.success("Ama desarquivada com sucesso!")
                        st.rerun()
        else:
            st.info("Ainda não existem amas registadas.")

# ==========================================
# 6. MÓDULO: FICHA DA CRIANÇA (COM CONSULTA E ARQUIVAMENTO)
# ==========================================
elif menu == "📄 Ficha da Criança":
    st.header("📄 Ficha da Criança (Inscrição, Consulta e Arquivamento)")
    st.caption("Decreto-Lei n.º 115/2015, de 22 de junho")
    
    aba_cri_1, aba_cri_2 = st.tabs(["✨ Nova Inscrição", "🔍 Consultar e Arquivar"])
    
    with aba_cri_1:
        with st.form("form_inscricao_crianca"):
            st.subheader("1. Atribuição de Ama Responsável")
            ama_responsavel_sel = st.selectbox("Ama Responsável", AMAS) if AMAS else st.text_input("Ama Responsável (Escreva o nome)")
            
            st.subheader("2. Dados Administrativos")
            col_adm1, col_adm2, col_adm3 = st.columns(3)
            processo = col_adm1.text_input("N.º de Processo")
            data_entrada = col_adm2.date_input("Data de Entrada", datetime.today())
            despacho = col_adm3.selectbox("Despacho", ["Pendente", "Aprovado", "Arquivado"])
            
            st.subheader("3. Identificação da Criança")
            c1, c2, c3, c4 = st.columns(4)
            nome_completo = c1.text_input("Nome Completo")
            naturalidade = c2.text_input("Naturalidade")
            data_nascimento = c3.date_input("Data de Nascimento", datetime(2023, 1, 1))
            sexo = c4.selectbox("Sexo", ["Masculino", "Feminino", "Outro"])
            
            c5, c6, c7, c8 = st.columns(4)
            num_identificacao = c5.text_input("N.º de Identificação (CC)")
            validade_identificacao = c6.text_input("Validade CC")
            nif = c7.text_input("NIF")
            niss = c8.text_input("NISS")
            
            st.subheader("4. Filiação")
            p1, p2 = st.columns(2)
            pai_nome = p1.text_input("Nome do Pai")
            pai_telemovel = p2.text_input("Telemóvel do Pai")
            
            m1, m2 = st.columns(2)
            mae_nome = m1.text_input("Nome da Mãe")
            mae_telemovel = m2.text_input("Telemóvel da Mãe")

            st.subheader("5. Anexos e Documentos")
            foto_file = st.file_uploader("Fotografia da Criança", type=['png', 'jpg', 'jpeg'], key="foto_crianca_up")
            docs_files = st.file_uploader("Anexar Documentos (PDF, etc.)", type=['pdf', 'png', 'jpg', 'jpeg', 'doc', 'docx'], accept_multiple_files=True, key="docs_crianca_up")

            submit_crianca = st.form_submit_button("💾 Guardar Ficha da Criança")

        # 👉 O IF FICA FORA DO WITH ST.FORM (alinhado com o 'with st.form')
        if submit_crianca:
            if nome_completo:
                # Conversão da Fotografia para Base64
                foto_b64 = ""
                if foto_file is not None:
                    foto_bytes = foto_file.read()
                    foto_b64 = base64.b64encode(foto_bytes).decode('utf-8')

                # Conversão dos Documentos Anexos para Base64
                docs_saved = []
                if docs_files:
                    for doc in docs_files:
                        doc_bytes = doc.read()
                        doc_b64 = base64.b64encode(doc_bytes).decode('utf-8')
                        docs_saved.append(f"{doc.name}:{doc_b64}")
                docs_str = "||".join(docs_saved)

                conn = get_db_connection()
                c = conn.cursor()
                c.execute('''INSERT INTO ficha_crianca 
                           (processo, data_entrada, despacho, ama_responsavel, nome_completo, naturalidade, data_nascimento, sexo,
                            num_identificacao, validade_identificacao, nif, niss, pai_nome, pai_telemovel, mae_nome, mae_telemovel, foto_path, docs_paths, arquivada)
                           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s, 0)''',
                          (processo, str(data_entrada), despacho, ama_responsavel_sel, nome_completo, naturalidade, str(data_nascimento), sexo,
                           num_identificacao, validade_identificacao, nif, niss, pai_nome, pai_telemovel, mae_nome, mae_telemovel, foto_b64, docs_str))
                conn.commit()
                c.close()
                conn.close()
                st.cache_data.clear()
                st.success(f"Ficha de {nome_completo} guardada com sucesso!")
                st.rerun()
            else:
                st.error("O Nome Completo da Criança é obrigatório.")
                
    with aba_cri_2:
        st.subheader("🔍 Consultar e Gerir Fichas de Crianças")
        
        # Atualizado para passar o role e o username de forma segura
        todas_criancas = get_criancas(
            incluir_arquivadas=True, 
            role=st.session_state.get('role', ''), 
            username=st.session_state.get('username', '')
        )
        
        if todas_criancas:
            crianca_selecionada = st.selectbox("Selecione a Criança para consultar:", todas_criancas, key="sel_cons_cri")
            # ... resto do código ...
            conn = get_db_connection()
            df_detalhe_cri = pd.read_sql_query("SELECT * FROM ficha_crianca WHERE nome_completo = %s", conn, params=(crianca_selecionada,))
            conn.close()
            
            if not df_detalhe_cri.empty:
                reg_cri = df_detalhe_cri.iloc[0]
                status_arq_c = "🔴 Arquivada" if reg_cri['arquivada'] == 1 else "🟢 Ativa"
                
                # Cabeçalho limpo com o estado
                col_hc1, col_hc2 = st.columns([3, 1])
                col_hc1.markdown(f"### Ficha do Educando: **{reg_cri['nome_completo']}**")
                col_hc2.markdown(f"**Estado:** {status_arq_c}")
                
                st.markdown("---")
                
                # Exibição organizada por colunas
                col_dc1, col_dc2 = st.columns(2)
                with col_dc1:
                    # Exibir Fotografia se existir
                    if reg_cri.get('foto_path') and len(reg_cri['foto_path']) > 50:
                        try:
                            foto_bytes_dec = base64.b64decode(reg_cri['foto_path'])
                            st.image(foto_bytes_dec, width=150, caption="Fotografia da Criança")
                        except Exception:
                            pass

                    st.markdown("#### 📄 Dados Administrativos")
                    st.write(f"**N.º de Processo:** {reg_cri['processo']}")
                    st.write(f"**Data de Entrada:** {reg_cri['data_entrada']}")
                    st.write(f"**Despacho / Estado:** {reg_cri['despacho']}")
                    st.write(f"**Ama Responsável:** {reg_cri['ama_responsavel']}")
                    st.markdown("#### 👦 Identificação da Criança")
                    st.write(f"**Naturalidade:** {reg_cri['naturalidade']}")
                    st.write(f"**Data de Nascimento:** {reg_cri['data_nascimento']}")
                    st.write(f"**Sexo:** {reg_cri['sexo']}")
                    st.write(f"**N.º Identificação (CC):** {reg_cri['num_identificacao']}")
                    st.write(f"**NIF / NISS:** {reg_cri['nif']} / {reg_cri['niss']}")
                
                with col_dc2:
                    st.markdown("#### 👪 Filiação e Contactos")
                    st.write(f"**Nome do Pai:** {reg_cri['pai_nome']}")
                    st.write(f"**Telemóvel do Pai:** {reg_cri['pai_telemovel']}")
                    st.write(f"**Nome da Mãe:** {reg_cri['mae_nome']}")
                    st.write(f"**Telemóvel da Mãe:** {reg_cri['mae_telemovel']}")

                    # Exibir Documentos Anexos se existirem
                    if reg_cri.get('docs_paths') and len(reg_cri['docs_paths']) > 5:
                        st.markdown("#### 📁 Documentos Anexos")
                        docs_list = reg_cri['docs_paths'].split("||")
                        for doc_item in docs_list:
                            if ":" in doc_item:
                                doc_name, doc_data_b64 = doc_item.split(":", 1)
                                try:
                                    doc_bytes_dec = base64.b64decode(doc_data_b64)
                                    st.download_button(
                                        label=f"📥 Descarregar {doc_name}",
                                        data=doc_bytes_dec,
                                        file_name=doc_name,
                                        key=f"dl_doc_{reg_cri['id']}_{doc_name}"
                                    )
                                except Exception:
                                    pass
                
                st.markdown("---")
                col_arq_c1, col_arq_c2 = st.columns(2)
                if reg_cri['arquivada'] == 0:
                    if col_arq_c1.button("📁 Arquivar Criança", key="btn_arq_cri"):
                        conn = get_db_connection()
                        c = conn.cursor()
                        c.execute("UPDATE ficha_crianca SET arquivada = 1 WHERE id = %s", (int(reg_cri['id']),))
                        conn.commit()
                        c.close()
                        conn.close()
                        st.cache_data.clear()
                        st.success("Criança arquivada com sucesso!")
                        st.rerun()
                else:
                    if col_arq_c2.button("📂 Desarquivar Criança", key="btn_desarq_cri"):
                        conn = get_db_connection()
                        c = conn.cursor()
                        c.execute("UPDATE ficha_crianca SET arquivada = 0 WHERE id = %s", (int(reg_cri['id']),))
                        conn.commit()
                        c.close()
                        conn.close()
                        st.cache_data.clear()
                        st.success("Criança desarquivada com sucesso!")
                        st.rerun()
        else:
            st.info("Ainda não existem crianças registadas.")

# ==========================================
# 7. MÓDULO: ROTINAS E CUIDADOS DIÁRIOS (AMA)
# ==========================================
elif menu == "📝 Rotinas e Cuidados Diários (Ama)":
    st.header("📝 Registo Diário de Rotinas e Cuidados")
    
    ama_atual = st.session_state.nome_associado if st.session_state.role == "Ama" else None
    criancas_disponiveis = get_criancas(ama_atual)
    
    with st.form("form_cuidados"):
        col1, col2, col3 = st.columns(3)
        data_registo = col1.date_input("Data do Registo", datetime.today())
        crianca = col2.selectbox("Criança", criancas_disponiveis if criancas_disponiveis else CRIANCAS)
        ama = col3.selectbox("Ama Responsável", AMAS, index=AMAS.index(ama_atual) if ama_atual in AMAS else 0)
        
        st.subheader("1. Bem-Estar")
        bem_estar = st.radio("Estado Geral / Disposição:", ["Bem-disposto", "Sensível", "Agitado", "Doente"], horizontal=True)
        
        st.subheader("2. Alimentação")
        col_a1, col_a2 = st.columns(2)
        almoco = col_a1.radio("Almoço:", ["Bem", "Mal"], horizontal=True)
        lanche = col_a2.radio("Lanche:", ["Bem", "Mal"], horizontal=True)
        
        st.subheader("3. Sono")
        col_s1, col_s2 = st.columns(2)
        sono_manha = col_s1.text_input("Sono da Manhã")
        sono_tarde = col_s2.text_input("Sono da Tarde")
        
        st.subheader("4. Higiene")
        col_h1, col_h2 = st.columns(2)
        fraldas = col_h1.number_input("Fraldas (N.º)", min_value=0, max_value=10, value=0)
        bacio = col_h2.number_input("Bacio (N.º)", min_value=0, max_value=10, value=0)
        
        st.subheader("5. Saúde e Recados")
        saude_medicacao = st.text_area("Saúde / Medicação")
        recados = st.text_input("Recados para os Pais")
        
        submit_cuidados = st.form_submit_button("💾 Guardar Registo de Rotinas")
        
        if submit_cuidados:
            conn = get_db_connection()
            c = conn.cursor()
            c.execute('''INSERT INTO registos_diarios 
                         (data, crianca, ama, bem_estar, almoco, lanche, sono_manha, sono_tarde, fraldas, bacio, saude_medicacao, recados)
                         VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)''',
                      (str(data_registo), crianca, ama, bem_estar, almoco, lanche, sono_manha, sono_tarde, fraldas, bacio, saude_medicacao, recados))
            conn.commit()
            c.close()
            conn.close()
            st.cache_data.clear()
            st.success(f"Registo de rotinas de {crianca} guardado com sucesso!")

# ==========================================
# 8. MÓDULO: ATIVIDADES E OCORRÊNCIAS
# ==========================================
elif menu == "🎨 Atividades e Ocorrências":
    st.header("🎨 Atividades, Desenvolvimento e Ocorrências")
    
    ama_atual = st.session_state.nome_associado if st.session_state.role == "Ama" else None
    criancas_disponiveis = get_criancas(ama_atual)
    
    with st.form("form_atividades"):
        col1, col2 = st.columns(2)
        data_registo = col1.date_input("Data", datetime.today())
        crianca = col2.selectbox("Criança", criancas_disponiveis if criancas_disponiveis else CRIANCAS)
        
        st.subheader("🧩 Atividades Sociopedagógicas")
        atividades = st.text_area("Descrição das brincadeiras")
        
        st.subheader("📈 Marcos de Aprendizagem")
        marcos = st.text_area("Novas conquistas")
        
        st.subheader("⚠️ Ocorrências / Incidentes")
        ocorrencias = st.text_area("Ocorrências")
        
        st.subheader("📌 Observações")
        observacoes = st.text_area("Observações")
        
        st.subheader("📸 Fotografias")
        fotos_files = st.file_uploader("Anexar Fotografias", type=['png', 'jpg', 'jpeg'], accept_multiple_files=True)
        
        submit_atividades = st.form_submit_button("💾 Guardar Atividades")
        
        if submit_atividades:
            fotos_saved = []
            if fotos_files:
                pasta_fotos = os.path.join("uploads", "atividades_fotos")
                os.makedirs(pasta_fotos, exist_ok=True)
                for foto in fotos_files:
                    foto_path = os.path.join(pasta_fotos, f"{crianca.replace(' ', '_')}_{datetime.today().strftime('%Y%m%d_%H%M%S')}_{foto.name}")
                    with open(foto_path, "wb") as f:
                        f.write(foto.getbuffer())
                    fotos_saved.append(foto_path)
            fotos_str = ", ".join(fotos_saved)
            
            conn = get_db_connection()
            c = conn.cursor()
            c.execute('''INSERT INTO atividades (data, crianca, atividades_sociopedagogicas, marcos_aprendizagem, ocorrencias, observacoes, fotos_paths)
                         VALUES (%s, %s, %s, %s, %s, %s, %s)''',
                      (str(data_registo), crianca, atividades, marcos, ocorrencias, observacoes, fotos_str))
            conn.commit()
            c.close()
            conn.close()
            st.cache_data.clear()
            st.success(f"Atividades registadas para {crianca}.")

# ==========================================
# 9. MÓDULO: SECRETARIA E ASSIDUIDADE (EXCEL BONITO)
# ==========================================
elif menu == "🕒 Secretaria e Assiduidade":
    st.header("🕒 Controlo de Frequência e Presenças")
    
    ama_atual = st.session_state.nome_associado if st.session_state.role == "Ama" else None
    criancas_disponiveis = get_criancas(ama_atual)
    
    with st.form("form_assiduidade"):
        col1, col2 = st.columns(2)
        data_registo = col1.date_input("Data", datetime.today())
        crianca = col2.selectbox("Criança", criancas_disponiveis if criancas_disponiveis else CRIANCAS)
        
        col3, col4 = st.columns(2)
        hora_entrada = col3.time_input("Hora de Entrada")
        hora_saida = col4.time_input("Hora de Saída", value=None)
        
        notas_assiduidade = st.text_input("Notas (ex: entregue pelo irmão)")
        
        submit_assiduidade = st.form_submit_button("💾 Registar Presença Manual")
        
        if submit_assiduidade:
            conn = get_db_connection()
            c = conn.cursor()
            c.execute('''INSERT INTO assiduidade (data, crianca, hora_entrada, hora_saida, notas)
                         VALUES (%s, %s, %s, %s, %s)''',
                      (str(data_registo), crianca, str(hora_entrada), str(hora_saida) if hora_saida else "", notas_assiduidade))
            conn.commit()
            c.close()
            conn.close()
            st.cache_data.clear()
            st.success("Registo de presença guardado com sucesso!")
            
    st.divider()
    st.subheader("📥 Exportar Mapa de Presenças Mensal em Excel (.xlsx)")
    
    col_ex1, col_ex2, col_ex3 = st.columns(3)
    crianca_excel = col_ex1.selectbox("Criança para o Mapa", criancas_disponiveis if criancas_disponiveis else CRIANCAS, key="excel_crianca")
    ama_excel = col_ex2.selectbox("Ama Responsável", AMAS if AMAS else ["Geral"])
    
    meses_dict = {
        "Janeiro": "01", "Fevereiro": "02", "Março": "03", "Abril": "04",
        "Maio": "05", "Junho": "06", "Julho": "07", "Agosto": "08",
        "Setembro": "09", "Outubro": "10", "Novembro": "11", "Dezembro": "12"
    }
    col_m1, col_m2 = col_ex3.columns(2)
    mes_nome = col_m1.selectbox("Mês", list(meses_dict.keys()), index=datetime.today().month - 1)
    ano_sel = col_m2.number_input("Ano", min_value=2024, max_value=2030, value=datetime.today().year)
    
    mes_num = meses_dict[mes_nome]
    ano_mes_str = f"{ano_sel}-{mes_num}"
    
    if st.button("📊 Gerar Excel do Mês Selecionado", type="primary"):
        conn = get_db_connection()
        query_excel = "SELECT data, hora_entrada, hora_saida, notas FROM assiduidade WHERE crianca = %s AND data LIKE %s ORDER BY data ASC"
        df_assiduidade = pd.read_sql_query(query_excel, conn, params=(crianca_excel, f"{ano_mes_str}%"))
        conn.close()
        
        if df_assiduidade.empty:
            st.warning(f"Não existem registos para {crianca_excel} em {mes_nome} de {ano_sel}.")
        else:
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df_assiduidade.to_excel(writer, index=False, sheet_name='Presenças')
                
                # Formatação bonita com openpyxl (cabeçalho azul, texto branco, larguras ajustadas)
                workbook = writer.book
                worksheet = writer.sheets['Presenças']
                
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                
                header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                align_center = Alignment(horizontal="center", vertical="center")
                
                for col_num in range(1, len(df_assiduidade.columns) + 1):
                    cell = worksheet.cell(row=1, column=col_num)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = align_center
                
                # Ajuste automático das colunas
                for col in worksheet.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = col[0].column_letter
                    worksheet.column_dimensions[col_letter].width = max(max_len + 4, 15)

            excel_data = output.getvalue()
            st.download_button(
                label=f"📥 Descarregar Mapa Formatado ({mes_nome} {ano_sel}) (.xlsx)",
                data=excel_data,
                file_name=f"Mapa_Presencas_{crianca_excel.replace(' ', '_')}_{ano_sel}_{mes_num}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )

# ==========================================
# 10. MÓDULO: ESPAÇO DOS PAIS
# ==========================================
elif menu == "👨‍👩‍👧 Espaço dos Pais (Rotinas e Presenças)":
    st.header("👨‍👩‍👧 Espaço dos Pais - Acompanhamento Diário")
    
    crianca_pai = st.session_state.nome_associado if st.session_state.role == "Pais" else st.selectbox("Selecione a Criança:", CRIANCAS)
    
    st.markdown(f"---")
    st.subheader(f"📌 Informações do Educando: {crianca_pai}")
    
    tab_p1, tab_p2, tab_p3 = st.tabs(["🕒 Entradas e Saídas", "📝 Rotinas e Cuidados", "🎨 Atividades e Fotos"])
    
    with tab_p1:
        st.markdown("### Registo de Presenças")
        conn = get_db_connection()
        df_pais_ass = pd.read_sql_query("SELECT data, hora_entrada, hora_saida, notas FROM assiduidade WHERE crianca = %s ORDER BY data DESC", conn, params=(crianca_pai,))
        conn.close()
        if df_pais_ass.empty:
            st.info("Ainda sem registos de presenças.")
        else:
            st.dataframe(df_pais_ass, use_container_width=True)
                
    with tab_p2:
        st.markdown("### Rotinas Diárias")
        conn = get_db_connection()
        df_pais_rot = pd.read_sql_query("SELECT data, ama, bem_estar, almoco, lanche, sono_manha, sono_tarde, fraldas, bacio, saude_medicacao, recados FROM registos_diarios WHERE crianca = %s ORDER BY data DESC", conn, params=(crianca_pai,))
        conn.close()
        if df_pais_rot.empty:
            st.info("Ainda sem registos de rotinas diárias.")
        else:
            for idx, row in df_pais_rot.iterrows():
                with st.expander(f"📅 Data: {row['data']} (Ama: {row['ama']})"):
                    st.write(f"**Disposição:** {row['bem_estar']}")
                    st.write(f"**Alimentação:** Almoço: {row['almoco']} | Lanche: {row['lanche']}")
                    st.write(f"**Sono:** Manhã: {row['sono_manha']} | Tarde: {row['sono_tarde']}")
                    st.write(f"**Higiene:** Fraldas: {row['fraldas']} | Bacio: {row['bacio']}")
                    if row['saude_medicacao']:
                        st.warning(f"**Medicação:** {row['saude_medicacao']}")
                    if row['recados']:
                        st.info(f"**Recados:** {row['recados']}")
                            
    with tab_p3:
        st.markdown("### Atividades e Fotografias")
        conn = get_db_connection()
        df_pais_ativ = pd.read_sql_query("SELECT data, atividades_sociopedagogicas, marcos_aprendizagem, fotos_paths FROM atividades WHERE crianca = %s ORDER BY data DESC", conn, params=(crianca_pai,))
        conn.close()
        if df_pais_ativ.empty:
            st.info("Ainda sem atividades ou fotografias registadas.")
        else:
            for idx, row in df_pais_ativ.iterrows():
                st.markdown(f"**📅 Data:** {row['data']}")
                st.markdown(f"**🧩 Atividades:** {row['atividades_sociopedagogicas']}")
                st.markdown(f"**📈 Marcos:** {row['marcos_aprendizagem']}")
                if row['fotos_paths']:
                    fotos_list = row['fotos_paths'].split(", ")
                    st.markdown("**📸 Fotografias:**")
                    cols = st.columns(len(fotos_list))
                    for i, fpath in enumerate(fotos_list):
                        if os.path.exists(fpath):
                            cols[i].image(fpath, width=150)
                st.markdown("---")

# ==========================================
# 11. MÓDULO: RELATÓRIOS E IMPRESSÃO
# ==========================================
elif menu == "📊 Relatórios e Impressão":
    st.header("📊 Relatórios de Atividades e Desenvolvimento")
    
    col_r1, col_r2, col_r3 = st.columns(3)
    crianca_rel = col_r1.selectbox("Selecionar Criança", CRIANCAS, key="rel_crianca")
    tipo_relatorio = col_r2.selectbox("Tipo de Relatório", ["Diário", "Semanal", "Mensal"])
    data_rel = col_r3.date_input("Data de Referência", datetime.today(), key="rel_data")
    ama_rel = st.selectbox("Ama Responsável", AMAS, key="rel_ama")
    
    if st.button("🔍 Gerar e Guardar Relatório", type="primary"):
        conn = get_db_connection()
        if tipo_relatorio == "Diário":
            data_str = str(data_rel)
            df_ativ = pd.read_sql_query("SELECT * FROM atividades WHERE crianca = %s AND data = %s", conn, params=(crianca_rel, data_str))
            periodo_txt = f"Dia: {data_rel.strftime('%d/%m/%Y')}"
        elif tipo_relatorio == "Semanal":
            inicio_semana = data_rel - timedelta(days=data_rel.weekday())
            fim_semana = inicio_semana + timedelta(days=6)
            df_ativ = pd.read_sql_query("SELECT * FROM atividades WHERE crianca = %s AND data BETWEEN %s AND %s ORDER BY data DESC", conn, params=(crianca_rel, str(inicio_semana), str(fim_semana)))
            periodo_txt = f"Semana de {inicio_semana.strftime('%d/%m/%Y')} a {fim_semana.strftime('%d/%m/%Y')}"
        else:
            ano_mes = data_rel.strftime('%Y-%m')
            df_ativ = pd.read_sql_query("SELECT * FROM atividades WHERE crianca = %s AND data LIKE %s ORDER BY data DESC", conn, params=(crianca_rel, f"{ano_mes}%"))
            periodo_txt = f"Mês: {data_rel.strftime('%B de %Y')}"
        conn.close()
            
        if df_ativ.empty:
            st.info("Não existem registos de atividades para o período selecionado.")
        else:
            doc = Document()
            p_title = doc.add_paragraph()
            r_title = p_title.add_run("RELATÓRIO DE ATIVIDADES E DESENVOLVIMENTO")
            r_title.font.name = 'Arial'
            r_title.font.size = Pt(14)
            r_title.font.bold = True
            r_title.font.color.rgb = RGBColor(31, 78, 121)
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            doc.add_paragraph(f"Criança: {crianca_rel}")
            doc.add_paragraph(f"Ama Responsável: {ama_rel}")
            doc.add_paragraph(f"Período: {periodo_txt}")
            doc.add_paragraph("-" * 50)
            
            for idx, row in df_ativ.iterrows():
                doc.add_paragraph(f"Data do Registo: {row['data']}")
                doc.add_paragraph(f"Atividades Sociopedagógicas:\n{row['atividades_sociopedagogicas']}")
                doc.add_paragraph(f"Marcos de Aprendizagem:\n{row['marcos_aprendizagem']}")
                doc.add_paragraph(f"Ocorrências / Incidentes:\n{row['ocorrencias'] if row['ocorrencias'] else 'Nenhuma ocorrência registada.'}")
                doc.add_paragraph(f"Observações:\n{row['observacoes'] if row['observacoes'] else 'Sem observações.'}")
                
                if row['fotos_paths']:
                    doc.add_paragraph("Registo Fotográfico:")
                    for fpath in row['fotos_paths'].split(", "):
                        if os.path.exists(fpath):
                            try:
                                doc.add_picture(fpath, width=Inches(3.0))
                            except Exception:
                                pass
                doc.add_paragraph("_" * 40)
            
            pasta_crianca = os.path.join("relatorios_criancas", crianca_rel.replace(' ', '_'))
            os.makedirs(pasta_crianca, exist_ok=True)
            nome_ficheiro_docx = f"Relatorio_{tipo_relatorio}_{datetime.today().strftime('%Y%m%d_%H%M%S')}.docx"
            caminho_completo = os.path.join(pasta_crianca, nome_ficheiro_docx)
            doc.save(caminho_completo)
            
            with open(caminho_completo, "rb") as f:
                docx_bytes = f.read()
                
            st.download_button(
                label="📥 Descarregar Relatório Formatado com Fotos (.docx)",
                data=docx_bytes,
                file_name=nome_ficheiro_docx,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary"
            )

# ==========================================
# 12. MÓDULO: SUPERVISÃO (EDUCADORA)
# ==========================================
elif menu == "👩‍🏫 Supervisão (Educadora)":
    st.header("👩‍🏫 Painel da Educadora de Infância")
    
    filtro_crianca = st.selectbox("Filtrar por Criança:", ["Todas"] + CRIANCAS)
    tab1, tab2, tab3 = st.tabs(["📋 Rotinas Diárias", "🎨 Atividades e Marcos", "⚠️ Ocorrências"])
    
    conn = get_db_connection()
    if filtro_crianca == "Todas":
        df_cuidados = pd.read_sql_query("SELECT data, crianca, ama, bem_estar, almoco, lanche, sono_manha, sono_tarde, fraldas, bacio, saude_medicacao FROM registos_diarios ORDER BY data DESC", conn)
        df_atividades = pd.read_sql_query("SELECT data, crianca, atividades_sociopedagogicas, marcos_aprendizagem, observacoes FROM atividades ORDER BY data DESC", conn)
        df_ocorrencias = pd.read_sql_query("SELECT data, crianca, ocorrencias FROM atividades WHERE ocorrencias != '' ORDER BY data DESC", conn)
    else:
        df_cuidados = pd.read_sql_query("SELECT data, crianca, ama, bem_estar, almoco, lanche, sono_manha, sono_tarde, fraldas, bacio, saude_medicacao FROM registos_diarios WHERE crianca = %s ORDER BY data DESC", conn, params=(filtro_crianca,))
        df_atividades = pd.read_sql_query("SELECT data, crianca, atividades_sociopedagogicas, marcos_aprendizagem, observacoes FROM atividades WHERE crianca = %s ORDER BY data DESC", conn, params=(filtro_crianca,))
        df_ocorrencias = pd.read_sql_query("SELECT data, crianca, ocorrencias FROM atividades WHERE crianca = %s AND ocorrencias != '' ORDER BY data DESC", conn, params=(filtro_crianca,))
    conn.close()
    
    with tab1:
        st.dataframe(df_cuidados, use_container_width=True)
    with tab2:
        st.dataframe(df_atividades, use_container_width=True)
    with tab3:
        if df_ocorrencias.empty:
            st.success("Sem ocorrências registadas.")
        else:
            st.error("Atenção: Ocorrências registadas que exigem leitura.")
            st.dataframe(df_ocorrencias, use_container_width=True)

# ==========================================
# 13. MÓDULO: GESTÃO DE UTILIZADORES
# ==========================================
elif menu == "👥 Gestão de Utilizadores":
    st.header("👥 Gestão de Utilizadores e Contas de Acesso")
    st.caption("Acesso reservado ao Administrador para criação de contas, redefinição de passwords e gestão de perfis.")
    
    tab_u1, tab_u2, tab_u3 = st.tabs(["➕ Criar Utilizador", "🔑 Redefinir Palavra-passe", "🗑️ Eliminar Utilizador"])
    
    with tab_u1:
        st.subheader("Criar Novo Utilizador")
        col_u1, col_u2 = st.columns(2)
        novo_username = col_u1.text_input("Nome de Utilizador (Username)", key="novo_user")
        novo_password = col_u2.text_input("Palavra-passe (Password)", type="password", key="novo_pass")
        
        col_u3, col_u4 = st.columns(2)
        novo_role = col_u3.selectbox("Perfil de Utilizador", ["Educadora", "Ama", "Pais", "Administrador"], key="novo_role_sel")
        
        if novo_role == "Ama":
            amas_list = get_amas()
            novo_associado = col_u4.selectbox("Ama Associada (Nome exato)", amas_list if amas_list else ["Nenhuma ama registada"], key="assoc_ama")
        elif novo_role == "Pais":
            criancas_list = get_criancas()
            novo_associado = col_u4.selectbox("Criança Associada (Nome exato do educando)", criancas_list if criancas_list else ["Nenhuma criança registada"], key="assoc_crianca")
        else:
            novo_associado = col_u4.text_input("Nome / Identificação", value="Direção", key="assoc_dir")
            
        if st.button("💾 Criar Utilizador", type="primary"):
            if novo_username and novo_password:
                try:
                    conn = get_db_connection()
                    c = conn.cursor()
                    c.execute("INSERT INTO utilizadores (username, password, role, nome_associado) VALUES (%s, %s, %s, %s)",
                              (novo_username, novo_password, novo_role, novo_associado))
                    conn.commit()
                    c.close()
                    conn.close()
                    st.cache_data.clear()
                    st.success(f"Utilizador `{novo_username}` criado com sucesso para o perfil **{novo_role}** ({novo_associado})!")
                    st.rerun()
                except Exception:
                    st.error(f"O nome de utilizador `{novo_username}` já existe ou ocorreu um erro.")
            else:
                st.error("O Nome de Utilizador e a Palavra-passe são obrigatórios.")

    with tab_u2:
        st.subheader("🔑 Alterar / Redefinir Palavra-passe")
        conn = get_db_connection()
        df_users_reset = pd.read_sql_query("SELECT username, role FROM utilizadores WHERE username != 'admin'", conn)
        conn.close()
        
        if not df_users_reset.empty:
            user_list_reset = df_users_reset['username'].tolist()
            user_to_reset = st.selectbox("Selecione a conta a alterar:", user_list_reset, key="reset_user_sel")
            nova_pass_reset = st.text_input("Nova Palavra-passe:", type="password", key="reset_pass_input")
            
            if st.button("🔄 Atualizar Palavra-passe", type="primary"):
                if nova_pass_reset:
                    conn = get_db_connection()
                    c = conn.cursor()
                    c.execute("UPDATE utilizadores SET password = %s WHERE username = %s", (nova_pass_reset, user_to_reset))
                    conn.commit()
                    c.close()
                    conn.close()
                    st.cache_data.clear()
                    st.success(f"Palavra-passe da conta `{user_to_reset}` alterada com sucesso!")
                else:
                    st.error("Por favor, digite a nova palavra-passe.")
        else:
            st.info("Não existem contas elegíveis para alterar a palavra-passe.")

    with tab_u3:
        st.subheader("🗑️ Eliminar Utilizador")
        conn = get_db_connection()
        df_users_del = pd.read_sql_query("SELECT username FROM utilizadores WHERE role != 'Administrador' AND username != 'admin'", conn)
        conn.close()
        user_list_del = df_users_del['username'].tolist() if not df_users_del.empty else []
        
        if user_list_del:
            user_to_delete = st.selectbox("Selecione o utilizador a eliminar", user_list_del, key="del_user_sel")
            if st.button("Eliminar Utilizador Selecionado", type="secondary"):
                conn = get_db_connection()
                c = conn.cursor()
                c.execute("DELETE FROM utilizadores WHERE username = %s", (user_to_delete,))
                conn.commit()
                c.close()
                conn.close()
                st.cache_data.clear()
                st.success(f"Utilizador `{user_to_delete}` eliminado com sucesso!")
                st.rerun()
        else:
            st.info("Não existem utilizadores secundários disponíveis para eliminação.")
                
    st.divider()
    st.subheader("Utilizadores Registados no Sistema")
    conn = get_db_connection()
    df_users = pd.read_sql_query("SELECT id, username, role, nome_associado FROM utilizadores", conn)
    conn.close()
    st.dataframe(df_users, use_container_width=True)

# ==========================================
# 14. MÓDULO: MURAL DE RECADOS E COMUNICAÇÃO (SEM DUPLICADOS)
# ==========================================
elif menu == "💬 Mural de Recados":
    st.header("💬 Mural de Recados e Comunicação")
    st.caption("Avisos institucionais, comunicação por grupo de ama e canal privado de mensagens.")
    
    perfil_atual = st.session_state.role
    associado_atual = st.session_state.nome_associado
    
    if perfil_atual in ["Administrador", "Educadora", "Ama"]:
        with st.form("form_novo_recado_prof"):
            st.subheader("📢 Publicar Recado no Mural")
            tipo_destino = st.selectbox("Destinatário", ["Geral (Todos)", "Grupo de Ama Específica", "Individual (Criança / Pais)"])
            
            alvo_destino = "Geral"
            if tipo_destino == "Grupo de Ama Específica":
                alvo_destino = st.selectbox("Selecionar Ama", AMAS)
            elif tipo_destino == "Individual (Criança / Pais)":
                alvo_destino = st.selectbox("Selecionar Criança", CRIANCAS)
                
            mensagem_texto = st.text_area("Mensagem / Recado")
            submit_recado = st.form_submit_button("📤 Publicar no Mural")
            
            if submit_recado:
                if mensagem_texto:
                    data_hora_atual = datetime.now().strftime("%d/%m/%Y às %H:%M")
                    remetente_nome = f"{st.session_state.username} ({st.session_state.role})"
                    
                    conn = get_db_connection()
                    c = conn.cursor()
                    c.execute("INSERT INTO recados (data_hora, remetente, perfil, tipo_destino, alvo_destino, mensagem, privado) VALUES (%s, %s, %s, %s, %s, %s, 0)",
                              (data_hora_atual, remetente_nome, st.session_state.role, tipo_destino, alvo_destino, mensagem_texto))
                    conn.commit()
                    c.close()
                    conn.close()
                    st.success("Recado publicado com sucesso!")
                    st.cache_data.clear()
                    st.rerun()
                else:
                    st.error("A mensagem não pode estar vazia.")
    else:
        with st.form("form_mensagem_privada_pai"):
            st.subheader("✉️ Enviar Mensagem Privada (Ama ou Educadora)")
            
            conn = get_db_connection()
            df_ama_filho = pd.read_sql_query("SELECT ama_responsavel FROM ficha_crianca WHERE nome_completo = %s", conn, params=(associado_atual,))
            conn.close()
            ama_do_filho = df_ama_filho.iloc[0]['ama_responsavel'] if not df_ama_filho.empty else "Ama Responsável"
            
            destinatario_privado = st.selectbox("Enviar para:", [f"Ama ({ama_do_filho})", "Educadora / Direção"])
            mensagem_privada = st.text_area("Escreva a sua mensagem privada:")
            
            submit_privado = st.form_submit_button("📤 Enviar Mensagem Privada")
            
            if submit_privado:
                if mensagem_privada:
                    data_hora_atual = datetime.now().strftime("%d/%m/%Y às %H:%M")
                    remetente_nome = f"Família de {associado_atual}"
                    alvo_alojado = ama_do_filho if "Ama" in destinatario_privado else "Direção"
                    
                    conn = get_db_connection()
                    c = conn.cursor()
                    c.execute("INSERT INTO recados (data_hora, remetente, perfil, tipo_destino, alvo_destino, mensagem, privado) VALUES (%s, %s, 'Pais', 'Privado', %s, %s, 1)",
                              (data_hora_atual, remetente_nome, alvo_alojado, mensagem_privada))
                    conn.commit()
                    c.close()
                    conn.close()
                    st.success("Mensagem privada enviada com sucesso!")
                    st.cache_data.clear()
                    st.rerun()
                else:
                    st.error("Escreva uma mensagem antes de enviar.")

    st.divider()
    st.subheader("📋 Mensagens e Mural")
    
    conn = get_db_connection()
    if perfil_atual in ["Administrador", "Educadora"]:
        df_recados = pd.read_sql_query("SELECT * FROM recados ORDER BY id DESC", conn)
    elif perfil_atual == "Ama":
        criancas_dela = get_criancas(associado_atual)
        if criancas_dela:
            query = "SELECT * FROM recados WHERE privado = 0 AND (tipo_destino = 'Geral (Todos)' OR (tipo_destino = 'Grupo de Ama Específica' AND alvo_destino = %s) OR (tipo_destino = 'Individual (Criança / Pais)' AND alvo_destino = ANY(%s))) OR (privado = 1 AND alvo_destino = %s) ORDER BY id DESC"
            df_recados = pd.read_sql_query(query, conn, params=(associado_atual, list(criancas_dela), associado_atual))
        else:
            query = "SELECT * FROM recados WHERE privado = 0 AND (tipo_destino = 'Geral (Todos)' OR (tipo_destino = 'Grupo de Ama Específica' AND alvo_destino = %s)) OR (privado = 1 AND alvo_destino = %s) ORDER BY id DESC"
            df_recados = pd.read_sql_query(query, conn, params=(associado_atual, associado_atual))
    else: # Pais
        df_ama_filho = pd.read_sql_query("SELECT ama_responsavel FROM ficha_crianca WHERE nome_completo = %s", conn, params=(associado_atual,))
        ama_do_filho = df_ama_filho.iloc[0]['ama_responsavel'] if not df_ama_filho.empty else ""
        
        query = "SELECT * FROM recados WHERE (privado = 0 AND (tipo_destino = 'Geral (Todos)' OR (tipo_destino = 'Grupo de Ama Específica' AND alvo_destino = %s) OR (tipo_destino = 'Individual (Criança / Pais)' AND alvo_destino = %s))) OR (privado = 1 AND remetente LIKE %s) ORDER BY id DESC"
        df_recados = pd.read_sql_query(query, conn, params=(ama_do_filho, associado_atual, f"%{associado_atual}%"))
    
    conn.close()

    # 👉 ADICIONE ESTA LINHA PARA FILTRAR QUALQUER DUPLICADO VISUALMENTE
    if not df_recados.empty:
        df_recados = df_recados.drop_duplicates(subset=['mensagem', 'remetente', 'data_hora'])
        
    if df_recados.empty:
        st.info("Não existem mensagens ou recados para exibir.")
    else:
        for idx, row in df_recados.iterrows():
            recado_id = row['id']
            is_privado = row['privado'] == 1
            
            if is_privado:
                st.markdown(f"🔒 **[Mensagem Privada] {row['remetente']}** — *{row['data_hora']}* (Para: {row['alvo_destino']})")
                st.warning(row['mensagem'])
            else:
                st.markdown(f"📢 **{row['remetente']}** — *{row['data_hora']}*")
                st.caption(f"Destino: **{row['tipo_destino']}** ({row['alvo_destino']})")
                st.info(row['mensagem'])
                
            conn = get_db_connection()
            df_respostas = pd.read_sql_query("SELECT DISTINCT * FROM respostas_recados WHERE recado_id = %s ORDER BY id ASC", conn, params=(recado_id,))
            conn.close()
            
            if not df_respostas.empty:
                for r_idx, r_row in df_respostas.iterrows():
                    st.markdown(f"&nbsp;&nbsp;&nbsp;&nbsp;💬 **{r_row['remetente']}** (*{r_row['data_hora']}*): {r_row['mensagem']}")
            
            if not is_privado:
                with st.form(f"form_resp_{recado_id}"):
                    resp_texto = st.text_input("Responder a este recado:", key=f"resp_input_{recado_id}")
                    submit_resp = st.form_submit_button("Responder")
                    if submit_resp:
                        if resp_texto:
                            resp_data_hora = datetime.now().strftime("%d/%m/%Y às %H:%M")
                            resp_remetente = f"{st.session_state.username} ({st.session_state.role})"
                            
                            conn = get_db_connection()
                            c = conn.cursor()
                            c.execute("INSERT INTO respostas_recados (recado_id, data_hora, remetente, mensagem) VALUES (%s, %s, %s, %s)",
                                      (recado_id, resp_data_hora, resp_remetente, resp_texto))
                            conn.commit()
                            c.close()
                            conn.close()
                            st.cache_data.clear()
                            st.success("Resposta enviada!")
                            st.rerun()
            st.markdown("---")
